"""生成 Binizox PDF/DOCX 解析与可引用检索的无敏感信息样例。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUTPUT_DIR = Path(__file__).resolve().parent
TITLE = "Binizox 个人资料解析测试文档"


def _set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def build_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        paragraph_properties = style._element.get_or_add_pPr()
        border = paragraph_properties.find(qn("w:pBdr"))
        if border is not None:
            paragraph_properties.remove(border)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    intro = document.add_paragraph(
        "本文件是 Binizox 的虚构测试资料，用于验证文档解析、页码定位、表格提取、版本管理和可见范围过滤。"
        "内容不对应任何真实企业制度。"
    )
    intro.paragraph_format.space_after = Pt(12)

    document.add_heading("1 申请材料", level=1)
    document.add_paragraph(
        "申请远程访问权限时，申请人需要提供姓名、所属部门、使用设备资产编号、访问系统名称、业务用途和预计使用期限。"
    )
    document.add_paragraph(
        "如申请管理员权限，还需要直属主管确认和信息安全复核。申请人不得在工单中填写口令、令牌或私钥。"
    )
    document.add_heading("1 1 材料清单", level=2)
    for item in (
        "申请人姓名和所属部门",
        "设备资产编号和操作系统",
        "访问系统名称及所需权限范围",
        "业务用途和预计使用期限",
        "直属主管确认",
    ):
        document.add_paragraph(item, style="List Bullet")

    document.add_page_break()
    document.add_heading("2 审批矩阵", level=1)
    document.add_paragraph(
        "本节故意放在第二页，供解析器验证跨页正文及来源定位。审批要求以当前活动版本为准。"
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Inches(1.55), Inches(2.2), Inches(2.85))
    for index, (text, width) in enumerate(zip(("权限类型", "必要确认", "附加材料"), widths, strict=True)):
        cell = table.rows[0].cells[index]
        cell.width = width
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "1F4E78")
        _set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    rows = (
        ("普通访问", "直属主管", "业务用途和期限"),
        ("管理员权限", "直属主管与信息安全", "权限范围和风险说明"),
        ("敏感数据访问", "数据负责人和信息安全", "数据分类及导出说明"),
    )
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column_index, (cell, text, width) in enumerate(zip(cells, values, widths, strict=True)):
            cell.width = width
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index % 2 == 0:
                _set_cell_shading(cell, "EAF2F8")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index < 2 else WD_ALIGN_PARAGRAPH.LEFT

    final_paragraph = document.add_paragraph(
        "提交后，系统应保留文档版本、页码和片段标识。新版本发布前，旧版本仍是唯一可检索的活动版本。"
    )
    final_paragraph.paragraph_format.space_before = Pt(12)
    document.save(path)


def build_pdf(path: Path) -> None:
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if not font_path.exists():
        raise FileNotFoundError(f"CJK font is unavailable: {font_path}")
    pdfmetrics.registerFont(TTFont("BinizoxCJK", str(font_path)))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BinizoxTitle",
        parent=styles["Title"],
        fontName="BinizoxCJK",
        fontSize=20,
        leading=27,
        alignment=TA_CENTER,
        textColor=colors.black,
        spaceAfter=18,
    )
    heading_style = ParagraphStyle(
        "BinizoxHeading",
        parent=styles["Heading1"],
        fontName="BinizoxCJK",
        fontSize=14,
        leading=20,
        textColor=colors.black,
        spaceBefore=10,
        spaceAfter=8,
    )
    body_style = ParagraphStyle(
        "BinizoxBody",
        parent=styles["BodyText"],
        fontName="BinizoxCJK",
        fontSize=10.5,
        leading=18,
        textColor=colors.black,
        spaceAfter=10,
    )
    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.85 * inch,
        rightMargin=0.85 * inch,
        topMargin=0.8 * inch,
        bottomMargin=0.8 * inch,
        title=TITLE,
        author="Binizox",
    )
    story = [
        Paragraph(TITLE, title_style),
        Paragraph(
            "本文件是 Binizox 的虚构测试资料，用于验证 PDF 逐页文本、表格和来源位置。内容不对应任何真实制度。",
            body_style,
        ),
        Paragraph("1 申请材料", heading_style),
        Paragraph(
            "申请远程访问权限时，需要提供姓名、所属部门、设备资产编号、访问系统名称、业务用途和预计使用期限。",
            body_style,
        ),
        Paragraph("管理员权限还需要直属主管确认和信息安全复核。不得提交口令、令牌或私钥。", body_style),
        PageBreak(),
        Paragraph("2 审批矩阵", heading_style),
        Paragraph("本节位于第二页，用于验证跨页正文及页码定位。", body_style),
    ]
    data = [
        ["权限类型", "必要确认", "附加材料"],
        ["普通访问", "直属主管", "业务用途和期限"],
        ["管理员权限", "直属主管与信息安全", "权限范围和风险说明"],
        ["敏感数据访问", "数据负责人和信息安全", "数据分类及导出说明"],
    ]
    table = Table(data, colWidths=(1.35 * inch, 2.0 * inch, 2.85 * inch), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "BinizoxCJK"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("LEADING", (0, 0), (-1, -1), 14),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#EAF2F8")),
                ("BACKGROUND", (0, 4), (-1, 4), colors.HexColor("#EAF2F8")),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.extend(
        [
            table,
            Spacer(1, 14),
            Paragraph("新版本发布前，旧版本仍是唯一可检索的活动版本。", body_style),
        ]
    )
    document.build(story)


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx(OUTPUT_DIR / "personal-reference-v1.docx")
    build_pdf(OUTPUT_DIR / "personal-reference-v1.pdf")
