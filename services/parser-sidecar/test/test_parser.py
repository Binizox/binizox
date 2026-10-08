from __future__ import annotations

from io import BytesIO

from app.parser import parse_document
from docx import Document


def _document_ref(format_: str) -> dict:
    return {
        "document_id": "policy",
        "tenant_id": "1",
        "storage_uri": "minio://knowledgebases/source",
        "format": format_,
        "version": "1",
        "knowledge_base_id": "kb-test",
        "filename": f"policy.{format_}",
        "title": "Remote Access Policy",
        "content_hash": "abc",
        "permission": {"access_level": "tenant", "user_uids": []},
    }


def test_docx_preserves_heading_page_and_table() -> None:
    document = Document()
    document.add_heading("Remote Access Policy", level=0)
    document.add_heading("Required Materials", level=1)
    document.add_paragraph("Provide the software name and business reason.")
    document.add_page_break()
    document.add_heading("Approval Matrix", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Access"
    table.cell(0, 1).text = "Approver"
    table.cell(1, 0).text = "Admin"
    table.cell(1, 1).text = "Security"
    stream = BytesIO()
    document.save(stream)

    parsed = parse_document(_document_ref("docx"), stream.getvalue())

    assert parsed["status"] == "succeeded"
    assert {block["location"]["page"] for block in parsed["blocks"]} == {1, 2}
    assert parsed["blocks"][2]["heading_path"] == [
        "Remote Access Policy",
        "Required Materials",
    ]
    assert parsed["tables"][0]["rows"][1] == ["Admin", "Security"]
    assert parsed["tables"][0]["location"]["page"] == 2


def test_corrupt_docx_returns_structured_failure() -> None:
    parsed = parse_document(_document_ref("docx"), b"not-a-docx")

    assert parsed["status"] == "failed"
    assert parsed["blocks"] == []
    assert parsed["tables"] == []
    assert parsed["errors"][0]["code"] == "parse_failed"
